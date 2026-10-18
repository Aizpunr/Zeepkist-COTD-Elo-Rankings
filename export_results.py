"""export_results.py: write the result workbooks into results/ (see results_store.py).

Run once to move the source of truth from the xlsx files to tracked JSON, or
again (after moving results/ aside) to re-import a hand-edited workbook. Then
it checks every sheet it wrote against the workbook cell for cell (dimensions,
values and value types).

  python export_results.py
"""
import os
import sys

import openpyxl

import results_store as rs

BOOKS = [
    # (id, xlsx, kind, only these sheets or None)
    ('cotd-1', 'Zeepkist COTDs 1-25.xlsx', 'cotd', None),
    ('cotd-26', 'Zeepkist COTDs 26-50.xlsx', 'cotd', None),
    ('cotd-51', 'Zeepkist COTDs 51-75.xlsx', 'cotd', None),
    ('cotd-76', 'COTDs 76-100.xlsx', 'cotd', None),
    ('cotd-101', 'COTDs 101-125.xlsx', 'cotd', None),
    ('cotd-126', 'COTD 126-130.xlsx', 'cotd', None),
    ('cotd-131', None, 'cotd', None),           # the current book, filename from elo_engine
    ('roulette', 'cup roulette.xlsx', 'roulette', None),
    ('troll', 'Troll cup.xlsx', 'troll', None),
    ('lexer', 'Zeepkist COTD Results (lexer 23-02).xlsx', 'lexer', ['Map Index']),
]


def sheet_cells(ws):
    return {(r, c): cell.value for (r, c), cell in ws._cells.items() if cell.value is not None}


def block_title(above):
    for row in above:
        v = row[0]
        if isinstance(v, str) and v.strip():
            return v.strip()
    return None


def export(base):
    root = os.path.join(base, rs.RESULTS_DIR)
    os.makedirs(os.path.join(root, 'cups'), exist_ok=True)
    manifest, used = [], set()
    current = _current_cotd_book(base)
    for book_id, xlsx, kind, only in BOOKS:
        xlsx = xlsx or current
        wb = openpyxl.load_workbook(os.path.join(base, xlsx), data_only=True)
        sheets_meta = []
        for sname in wb.sheetnames:
            if only and sname not in only:
                continue
            ws = wb[sname]
            idx = len(sheets_meta)
            sheets_meta.append({'name': sname, 'max_row': ws.max_row, 'max_column': ws.max_column})
            cells = sheet_cells(ws)
            heads = sorted((r, c) for (r, c), v in cells.items() if v == 'Position')
            owned = set()
            for hr, hc in heads:
                cols = range(hc, hc + rs.BLOCK_WIDTH)
                if any(c in owned for c in cols):
                    raise SystemExit(f'{xlsx} / {sname}: blocks overlap at column {hc}')
                owned.update(cols)
                last = max([r for (r, c) in cells if c in cols] + [hr])
                grid = lambda r: [cells.get((r, c)) for c in cols]
                above = [grid(r) for r in range(1, hr)]
                block = {
                    'book': book_id, 'sheet': sname, 'col': hc, 'header_row': hr,
                    'above': above,
                    'header': grid(hr),
                    'rows': [grid(r) for r in range(hr + 1, last + 1)],
                }
                title = block_title(above) or f'{book_id} {sname} col {hc}'
                slug = rs.slugify(title)
                if slug in used:
                    slug = rs.slugify(f'{title} also in {book_id}')
                if slug in used:
                    raise SystemExit(f'duplicate block slug {slug}')
                used.add(slug)
                with open(os.path.join(root, 'cups', slug + '.json'), 'w', encoding='utf-8', newline='\n') as f:
                    f.write(rs.dump_block(block))
            rest = sorted([r, c, v] for (r, c), v in cells.items() if c not in owned)
            if rest:
                p = rs.sheet_file(book_id, idx, base)
                os.makedirs(os.path.dirname(p), exist_ok=True)
                with open(p, 'w', encoding='utf-8', newline='\n') as f:
                    f.write(rs.dump_sheet({'cells': rest}))
        manifest.append({'id': book_id, 'xlsx': xlsx, 'kind': kind, 'sheets': sheets_meta})
        wb.close()
    with open(os.path.join(root, 'books.json'), 'w', encoding='utf-8', newline='\n') as f:
        f.write(rs.dump_books(manifest))
    return manifest


def _current_cotd_book(base):
    import re
    with open(os.path.join(base, 'elo_engine.py'), encoding='utf-8') as f:
        src = f.read()
    found = re.findall(r"'(COTD \d+-\d+\.xlsx)'", src)
    if found:
        return found[-1]
    return rs.book_files('cotd', base)[-1]


def audit(base):
    """Every exported sheet equals its workbook: same dims, same (type, value) per cell."""
    rs.clear_cache()
    bad = 0
    for b in rs.books(base):
        real = openpyxl.load_workbook(os.path.join(base, b['xlsx']), data_only=True)
        mine = rs.load_workbook(b['xlsx'], base=base)
        for sh in b['sheets']:
            ws, ms = real[sh['name']], mine[sh['name']]
            if (ws.max_row, ws.max_column) != (ms.max_row, ms.max_column):
                print(f'DIMS {b["xlsx"]} / {sh["name"]}'); bad += 1; continue
            for ra, rb in zip(ws.iter_rows(values_only=True), ms.iter_rows(values_only=True)):
                if [(type(x), x) for x in ra] != [(type(x), x) for x in rb]:
                    print(f'CELLS {b["xlsx"]} / {sh["name"]}'); bad += 1; break
        real.close()
    return bad


if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    base = rs.BASE
    root = os.path.join(base, rs.RESULTS_DIR)
    if os.path.exists(root):
        raise SystemExit('results/ exists; move it aside (Recycle Bin, or rename) to re-import')
    m = export(base)
    print(f'exported {len(m)} books, {sum(len(b["sheets"]) for b in m)} sheets, '
          f'{len(os.listdir(os.path.join(root, "cups")))} blocks')
    bad = audit(base)
    print('audit: every sheet matches its workbook' if not bad else f'audit: {bad} sheet(s) differ')
    sys.exit(1 if bad else 0)
