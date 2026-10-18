"""sync_xlsx.py: bring the xlsx backups in line with results/ (the source).

Writes every cell that differs, clears cells results/ no longer has, adds
missing sheets, and follows the current book's rename when a new cup pushes
it past its range (COTD 131-166.xlsx -> COTD 131-167.xlsx; the old file goes
to the Recycle Bin after the new one is saved). Each book it changes is first
copied to backups/. Books already in sync are not touched. Ends with the
cell-for-cell audit from export_results.py.

  python sync_xlsx.py [--dry-run]
"""
import datetime
import glob
import os
import re
import shutil
import subprocess
import sys

import openpyxl

import results_store as rs
from export_results import audit

BASE = rs.BASE


def recycle(path):
    """Send a file to the Windows Recycle Bin, never a permanent delete."""
    r = subprocess.run([
        'powershell', '-NoProfile', '-Command',
        'Add-Type -AssemblyName Microsoft.VisualBasic; '
        "[Microsoft.VisualBasic.FileIO.FileSystem]::DeleteFile("
        f"'{path}', 'OnlyErrorDialogs', 'SendToRecycleBin')"
    ], capture_output=True, text=True)
    return r.returncode == 0


def _locate(book, base):
    """(path to open or None, path to recycle after saving or None)."""
    path = os.path.join(base, book['xlsx'])
    if os.path.exists(path):
        return path, None
    m = re.fullmatch(r'COTD (\d+)-\d+\.xlsx', book['xlsx'])
    if m:
        olds = [p for p in glob.glob(os.path.join(base, f'COTD {m.group(1)}-*.xlsx'))
                if re.fullmatch(r'COTD \d+-\d+\.xlsx', os.path.basename(p))]
        if len(olds) == 1:
            return olds[0], olds[0]
        if len(olds) > 1:
            raise SystemExit(f'{book["xlsx"]}: several older books, refusing to guess: {olds}')
    return None, None


def sync_book(book, base, dry_run):
    src, old = _locate(book, base)
    if src is None and book['kind'] == 'lexer':
        print(f'  {book["xlsx"]}: not on disk, skipped (Lexer\'s workbook, only Map Index is kept)')
        return False
    wb = openpyxl.load_workbook(src) if src else openpyxl.Workbook()
    if not src:
        wb.remove(wb.active)
    want = rs.load_workbook(book['xlsx'], base=base)
    changes = 0
    for sh in book['sheets']:
        if sh['name'] in wb.sheetnames:
            ws = wb[sh['name']]
        else:
            ws = wb.create_sheet(sh['name'])
            changes += 1
        target = want[sh['name']]._cells
        have = {rc: c.value for rc, c in ws._cells.items() if c.value is not None}
        for rc in set(have) | set(target):
            a, b = have.get(rc), target.get(rc)
            if type(a) is not type(b) or a != b:
                cell = ws.cell(row=rc[0], column=rc[1])
                cell.value = b
                if isinstance(b, str) and b.startswith('='):
                    cell.data_type = 's'  # a name like '=XDC=Wolf', not a formula
                changes += 1
        if (ws.max_row, ws.max_column) != (sh['max_row'], sh['max_column']):
            # openpyxl only saves an empty cell that has a style, so the corner
            # gets a text format to keep the sheet's extent.
            ws.cell(row=sh['max_row'], column=sh['max_column']).number_format = '@'
            changes += 1
    renamed = old is not None
    if not changes and not renamed:
        return False
    what = f'{changes} cell(s)' + (f', renamed from {os.path.basename(old)}' if renamed else '')
    if dry_run:
        print(f'  {book["xlsx"]}: would write {what}')
        return False
    if src:
        os.makedirs(os.path.join(base, 'backups'), exist_ok=True)
        stamp = datetime.datetime.now().strftime('%Y%m%dT%H%M%S')
        shutil.copy2(src, os.path.join(base, 'backups', f'{os.path.basename(src)}.pre_sync_{stamp}.xlsx'))
    dest = os.path.join(base, book['xlsx'])
    wb.save(dest)
    print(f'  {book["xlsx"]}: wrote {what}')
    if renamed and os.path.exists(dest):
        if recycle(old):
            print(f'  {os.path.basename(old)} -> Recycle Bin')
        else:
            print(f'  WARNING: could not recycle {os.path.basename(old)}; left in place')
    return True


def main(base=BASE, dry_run=False):
    changed = [b['xlsx'] for b in rs.books(base) if sync_book(b, base, dry_run)]
    if dry_run:
        return 0
    print(f'{len(changed)} book(s) written' if changed else 'xlsx backups already in sync')
    bad = audit(base)
    print('audit: every sheet matches results/' if not bad else f'audit: {bad} sheet(s) differ')
    return 1 if bad else 0


if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    sys.exit(main(dry_run='--dry-run' in sys.argv))
